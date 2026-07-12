from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(r'D:\project-iot-smart-home')
MD_PATH = ROOT / 'docs' / 'bao-cao-cuoi-ky.md'
OUT_PATH = ROOT / 'docs' / 'bao-cao-cuoi-ky.docx'

CONTENT_WIDTH_DXA = 9360


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def set_cell_width(cell, width_dxa: int):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.first_child_found_in('w:tcW')
    if tcW is None:
        tcW = OxmlElement('w:tcW')
        tcPr.append(tcW)
    tcW.set(qn('w:w'), str(width_dxa))
    tcW.set(qn('w:type'), 'dxa')


def set_table_borders(table, color='DADCE0', size='4'):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = tblPr.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tblPr.append(borders)
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        tag = f'w:{edge}'
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), size)
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), color)


def set_table_width(table, width_dxa=CONTENT_WIDTH_DXA):
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.first_child_found_in('w:tblW')
    if tblW is None:
        tblW = OxmlElement('w:tblW')
        tblPr.append(tblW)
    tblW.set(qn('w:w'), str(width_dxa))
    tblW.set(qn('w:type'), 'dxa')
    tblInd = tblPr.first_child_found_in('w:tblInd')
    if tblInd is None:
        tblInd = OxmlElement('w:tblInd')
        tblPr.append(tblInd)
    tblInd.set(qn('w:w'), '0')
    tblInd.set(qn('w:type'), 'dxa')


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    trPr.append(tbl_header)


def set_paragraph_spacing(paragraph, before=0, after=8, line=1.15):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def add_runs_with_inline_code(paragraph, text: str, bold=False, italic=False):
    parts = re.split(r'(`[^`]+`|\*\*[^*]+\*\*)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Courier New'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Courier New')
            run.font.size = Pt(10)
        elif part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            run = paragraph.add_run(part)
            run.bold = bold
            run.italic = italic
        if not (part.startswith('`') and part.endswith('`')):
            run.font.name = 'Arial'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
        run.font.color.rgb = RGBColor(0, 0, 0)


def configure_styles(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Arial'
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, before, after, color in [
        ('Heading 1', 20, 20, 6, RGBColor(0, 0, 0)),
        ('Heading 2', 16, 18, 6, RGBColor(0, 0, 0)),
        ('Heading 3', 14, 16, 4, RGBColor(67, 67, 67)),
        ('Heading 4', 12, 12, 4, RGBColor(67, 67, 67)),
    ]:
        style = styles[name]
        style.font.name = 'Arial'
        style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
        style.font.size = Pt(size)
        style.font.bold = False
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15


def add_title(doc: Document, text: str):
    p = doc.add_paragraph()
    p.style = doc.styles['Normal']
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = 'Arial'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    run.font.size = Pt(26)
    run.font.bold = False
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_subtitle(doc: Document, text: str):
    p = doc.add_paragraph()
    p.style = doc.styles['Normal']
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(text)
    run.font.name = 'Arial'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor(85, 85, 85)


def add_code_block(doc: Document, lines: list[str]):
    for line in lines:
        p = doc.add_paragraph()
        p.style = doc.styles['Normal']
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(line if line else ' ')
        run.font.name = 'Courier New'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Courier New')
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0, 0, 0)
    if lines:
        doc.paragraphs[-1].paragraph_format.space_after = Pt(8)


def parse_table(lines: list[str]):
    rows = []
    for line in lines:
        cells = [c.strip() for c in line.strip().strip('|').split('|')]
        rows.append(cells)
    if len(rows) >= 2 and all(re.fullmatch(r':?-{3,}:?', c.replace(' ', '')) for c in rows[1]):
        rows.pop(1)
    return rows


def add_table(doc: Document, table_lines: list[str]):
    rows = parse_table(table_lines)
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_width(table)
    set_table_borders(table)

    # Width heuristic: shorter first/status columns, wider description columns.
    widths = [CONTENT_WIDTH_DXA // cols] * cols
    if cols == 2:
        widths = [2600, CONTENT_WIDTH_DXA - 2600]
    elif cols == 3:
        widths = [2200, 2600, CONTENT_WIDTH_DXA - 4800]
    elif cols == 4:
        widths = [2100, 1900, 1900, CONTENT_WIDTH_DXA - 5900]
    elif cols == 5:
        widths = [1700, 1700, 1700, 1700, CONTENT_WIDTH_DXA - 6800]

    for row_idx, row_data in enumerate(rows):
        row = table.rows[row_idx]
        if row_idx == 0:
            set_repeat_table_header(row)
        for col_idx in range(cols):
            cell = row.cells[col_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_width(cell, widths[col_idx])
            text = row_data[col_idx] if col_idx < len(row_data) else ''
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            add_runs_with_inline_code(p, text, bold=(row_idx == 0))
            for run in p.runs:
                run.font.size = Pt(9.5 if cols >= 4 else 10)
                if row_idx == 0:
                    run.bold = True
    doc.add_paragraph()


def add_bullet(doc: Document, text: str, level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent = Inches(0.5 + 0.25 * level)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    add_runs_with_inline_code(p, text)


def add_number(doc: Document, text: str, level=0):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.left_indent = Inches(0.5 + 0.25 * level)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    add_runs_with_inline_code(p, text)


def build_docx():
    md = MD_PATH.read_text(encoding='utf-8')
    lines = md.splitlines()
    doc = Document()
    configure_styles(doc)

    i = 0
    title_count = 0
    in_code = False
    code_lines: list[str] = []
    table_lines: list[str] = []

    def flush_table():
        nonlocal table_lines
        if table_lines:
            add_table(doc, table_lines)
            table_lines = []

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()

        if line.startswith('```'):
            flush_table()
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if line.strip().startswith('|') and line.strip().endswith('|'):
            table_lines.append(line)
            i += 1
            continue
        flush_table()

        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped == '---':
            i += 1
            continue

        heading = re.match(r'^(#{1,6})\s+(.*)$', stripped)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2).strip()
            if level == 1 and title_count == 0:
                add_title(doc, text)
                title_count += 1
            elif level == 1 and title_count == 1:
                add_subtitle(doc, text)
                title_count += 1
            else:
                style_level = min(level, 4)
                p = doc.add_paragraph(style=f'Heading {style_level}')
                add_runs_with_inline_code(p, text)
            i += 1
            continue

        bullet = re.match(r'^-\s+(.*)$', stripped)
        if bullet:
            add_bullet(doc, bullet.group(1))
            i += 1
            continue

        number = re.match(r'^\d+\.\s+(.*)$', stripped)
        if number:
            add_number(doc, number.group(1))
            i += 1
            continue

        # Placeholder figure text gets italic muted formatting.
        p = doc.add_paragraph()
        set_paragraph_spacing(p)
        if stripped.startswith('**[Placeholder') or stripped.startswith('**[Placeholder'.lower()):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            clean = stripped.replace('**', '')
            run = p.add_run(clean)
            run.italic = True
            run.font.name = 'Arial'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(85, 85, 85)
        else:
            add_runs_with_inline_code(p, stripped)
        i += 1

    flush_table()

    # Footer page number is intentionally omitted for Google Docs native feel.
    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == '__main__':
    build_docx()
